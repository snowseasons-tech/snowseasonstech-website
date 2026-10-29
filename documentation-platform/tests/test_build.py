import os
from pathlib import Path

import pytest
import yaml
from docx import Document

from generator import pdf_export
from generator.build import BuildError, build_product, check_product
from generator.registry import (
    RegistryError,
    list_product_ids,
    load_catalog,
    load_product,
    product_file,
    resolve_product,
)
from generator.renderer import make_environment, render_document
from generator.validation import validate_product, validate_rendered

ROOT = Path(__file__).resolve().parents[1]
PRODUCTS = ROOT / "products"
TEMPLATES = ROOT / "templates"
STEMS = ("product-overview", "getting-started", "pricing-sheet")
FAKE_CATALOG = {
    "bookscan": {
        "name": "BookScan Live",
        "stripeProduct": "prod_test",
        "version": "1.0.0",
        "blurb": "Catalog blurb.",
    }
}
SNOWSTORE = os.environ.get("SNOWSTORE_DIR")


def cdot() -> dict:
    return load_product(product_file(PRODUCTS, "cdot-stations"))


def test_registry_definitions_are_valid():
    for product_id in list_product_ids(PRODUCTS):
        if load_product(product_file(PRODUCTS, product_id))["product"].get("catalog_id"):
            continue  # needs the storefront catalog; see the snowStore test below
        assert check_product(product_id, PRODUCTS).ok, product_id


def test_rejects_unknown_status():
    data = cdot()
    data["product"]["status"] = "shipping-soon"
    assert any("status" in error for error in validate_product(data).errors)


def test_rejects_unsupported_document_type():
    data = cdot()
    data["product"]["documentation"]["generate"].append("user_manual")
    assert any("user_manual" in error for error in validate_product(data).errors)


def test_rejects_unquoted_version_that_is_not_semver():
    data = cdot()
    data["product"]["version"] = 1.0
    assert any("version" in error for error in validate_product(data).errors)


def test_warns_until_features_are_confirmed():
    report = validate_product(cdot())
    assert report.ok
    assert any("features_confirmed" in warning for warning in report.warnings)


def test_overview_lists_each_registered_feature_once():
    product = cdot()["product"]
    text = render_document(make_environment(TEMPLATES), "product_overview", product)
    for feature in product["features"]:
        assert text.count(f"- {feature}") == 1


def test_rendered_documents_have_no_unfilled_markers():
    product = cdot()["product"]
    env = make_environment(TEMPLATES)
    for doc_type in product["documentation"]["generate"]:
        assert validate_rendered(render_document(env, doc_type, product)) == []


def test_unfilled_marker_is_caught():
    assert validate_rendered("Price: TBD")


def test_build_writes_docx_and_manifest(tmp_path):
    manifest = build_product(
        "cdot-stations",
        products_dir=PRODUCTS,
        templates_dir=TEMPLATES,
        out_dir=tmp_path,
        formats=("docx",),
    )
    target = tmp_path / "cdot-stations"
    for stem in STEMS:
        assert (target / f"{stem}.md").is_file()
        assert (target / f"{stem}.docx").is_file()
    assert (target / "manifest.json").is_file()
    assert manifest["version"] == "0.1.0"
    assert len(manifest["outputs"]) == 6

    overview = Document(target / "product-overview.docx")
    assert "CDOT Stations" in overview.paragraphs[0].text
    assert len(Document(target / "pricing-sheet.docx").tables) == 1


def test_pdf_is_skipped_without_pandoc(tmp_path, monkeypatch):
    monkeypatch.setattr(pdf_export, "pandoc_available", lambda: False)
    manifest = build_product(
        "cdot-stations",
        products_dir=PRODUCTS,
        templates_dir=TEMPLATES,
        out_dir=tmp_path,
        formats=("pdf",),
    )
    assert {item["file"] for item in manifest["skipped"]} == {f"{stem}.pdf" for stem in STEMS}


def test_build_refuses_invalid_product(tmp_path):
    bad_dir = tmp_path / "products" / "bad-product"
    bad_dir.mkdir(parents=True)
    (bad_dir / "product.yaml").write_text(yaml.safe_dump({"product": {"id": "bad-product"}}))
    with pytest.raises(BuildError):
        build_product(
            "bad-product",
            products_dir=tmp_path / "products",
            templates_dir=TEMPLATES,
            out_dir=tmp_path / "build",
        )


def test_catalog_fills_name_version_and_blurb():
    entry = load_product(product_file(PRODUCTS, "bookscan"))
    product = resolve_product(entry, FAKE_CATALOG)["product"]
    assert product["name"] == "BookScan Live"
    assert product["version"] == "1.0.0"
    assert product["description"] == "Catalog blurb."
    assert product["billing"]["plans"][0]["stripe_product"] == "prod_test"
    assert validate_product(resolve_product(entry, FAKE_CATALOG)).ok


def test_docs_entry_overrides_catalog_fields():
    data = load_product(product_file(PRODUCTS, "bookscan"))
    data["product"]["description"] = "Docs description."
    product = resolve_product(data, FAKE_CATALOG)["product"]
    assert product["description"] == "Docs description."


def test_catalog_id_without_store_is_an_error():
    with pytest.raises(RegistryError, match="--store"):
        resolve_product(load_product(product_file(PRODUCTS, "bookscan")), None)


def test_unknown_catalog_id_is_an_error():
    data = load_product(product_file(PRODUCTS, "bookscan"))
    data["product"]["catalog_id"] = "not-a-product"
    with pytest.raises(RegistryError, match="not-a-product"):
        resolve_product(data, FAKE_CATALOG)


@pytest.mark.skipif(not SNOWSTORE, reason="set SNOWSTORE_DIR to test against a snowStore checkout")
def test_real_snowstore_catalog_resolves_every_catalog_entry():
    catalog = load_catalog(Path(SNOWSTORE))
    for product_id in list_product_ids(PRODUCTS):
        assert check_product(product_id, PRODUCTS, catalog).ok, product_id
