"""Convert the Markdown subset used by our templates into an editable .docx."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

HEADING = re.compile(r"^(#{1,6})\s+(.*)")
BULLET = re.compile(r"^\s*[-*]\s+(.*)")
NUMBERED = re.compile(r"^\s*\d+\.\s+(.*)")
BOLD = re.compile(r"\*\*(.+?)\*\*")
TABLE_DIVIDER = re.compile(r"^:?-+:?$")


def _add_runs(paragraph: Paragraph, text: str) -> None:
    text = text.replace("`", "")
    position = 0
    for match in BOLD.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        paragraph.add_run(match.group(1)).bold = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def _cells(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def _add_table(document: Document, rows: list[str]) -> None:
    header, *body = rows
    table = document.add_table(rows=0, cols=len(_cells(header)))
    table.style = "Table Grid"
    for index, line in enumerate(rows):
        cells = _cells(line)
        if all(TABLE_DIVIDER.match(cell) for cell in cells):
            continue
        row = table.add_row()
        for cell, text in zip(row.cells, cells):
            paragraph = cell.paragraphs[0]
            _add_runs(paragraph, text)
            if index == 0:
                for run in paragraph.runs:
                    run.bold = True


def markdown_to_docx(markdown: str, path: Path) -> None:
    document = Document()
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
        elif stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(lines[index])
                index += 1
            _add_table(document, rows)
        elif match := HEADING.match(stripped):
            document.add_heading(match.group(2), level=min(len(match.group(1)), 4))
            index += 1
        elif match := BULLET.match(line):
            _add_runs(document.add_paragraph(style="List Bullet"), match.group(1))
            index += 1
        elif match := NUMBERED.match(line):
            _add_runs(document.add_paragraph(style="List Number"), match.group(1))
            index += 1
        else:
            _add_runs(document.add_paragraph(), stripped)
            index += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
